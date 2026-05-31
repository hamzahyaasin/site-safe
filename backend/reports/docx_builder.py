from io import BytesIO

from django.db.models import Count
from docx import Document

from alerts.models import Alert
from workers.models import Worker


def _save_document(doc):
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_incident_report(date_from, date_to):
    doc = Document()
    doc.add_heading("Site-Safe Incident Report", level=0)
    doc.add_paragraph(f"Reporting period: {date_from} to {date_to}")

    alerts = (
        Alert.objects.filter(timestamp__date__gte=date_from, timestamp__date__lte=date_to)
        .select_related("worker")
        .order_by("-timestamp")
    )

    doc.add_paragraph(f"Total incidents: {alerts.count()}")

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Timestamp", "Type", "Severity", "Worker", "Source", "Status"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    for alert in alerts[:500]:
        row = table.add_row().cells
        row[0].text = alert.timestamp.strftime("%Y-%m-%d %H:%M")
        row[1].text = alert.alert_type
        row[2].text = alert.severity
        row[3].text = alert.worker.name if alert.worker_id else "—"
        row[4].text = alert.source
        row[5].text = "Resolved" if alert.is_resolved else "Open"

    return _save_document(doc)


def build_compliance_report(date_from, date_to):
    doc = Document()
    doc.add_heading("Site-Safe Compliance Report", level=0)
    doc.add_paragraph(f"Reporting period: {date_from} to {date_to}")

    total_workers = Worker.objects.filter(is_active=True).count()
    alerts_qs = Alert.objects.filter(timestamp__date__gte=date_from, timestamp__date__lte=date_to)
    total_alerts = alerts_qs.count()
    resolved = alerts_qs.filter(is_resolved=True).count()
    resolution_rate = round((resolved / total_alerts) * 100, 1) if total_alerts else 100.0

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(f"Active workers: {total_workers}")
    doc.add_paragraph(f"Alerts in period: {total_alerts}")
    doc.add_paragraph(f"Resolved alerts: {resolved}")
    doc.add_paragraph(f"Resolution rate: {resolution_rate}%")

    doc.add_heading("Alerts by Type", level=1)
    type_table = doc.add_table(rows=1, cols=2)
    type_table.style = "Table Grid"
    type_table.rows[0].cells[0].text = "Alert Type"
    type_table.rows[0].cells[1].text = "Count"

    counts = alerts_qs.values("alert_type").annotate(c=Count("id")).order_by("-c")
    for row in counts:
        cells = type_table.add_row().cells
        cells[0].text = row["alert_type"]
        cells[1].text = str(row["c"])

    ppe_violations = alerts_qs.filter(alert_type="PPE_VIOLATION").count()
    doc.add_paragraph(f"PPE violations: {ppe_violations}")

    return _save_document(doc)


def build_report_document(report_type, date_from, date_to):
    if report_type == "INCIDENT":
        return build_incident_report(date_from, date_to)
    return build_compliance_report(date_from, date_to)
