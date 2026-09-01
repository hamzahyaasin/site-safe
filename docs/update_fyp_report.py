#!/usr/bin/env python3
"""Build and finalize the rubric-aligned Site-Safe FYP report.

The build pass makes the content and structural edits and writes front-matter
page-number placeholders. The finalize pass reads a rendered PDF, measures the
actual heading/caption pages, and writes those measured numbers into the static
table of contents and lists of figures/tables.
"""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from pypdf import PdfReader


CHAPTER_TITLES = {
    1: "Introduction",
    2: "Review of Existing Systems",
    3: "Proposed System Design",
    4: "Hardware and Software Implementation",
    5: "Testing, Results and Discussion",
    6: "Conclusion and Future Work",
}

TABLE_41 = [
    ("Tool / Platform", "Version / Status", "Purpose"),
    ("Python", "3.14.3", "Backend logic and model utilities; measured in the project virtual environment"),
    ("Django", "6.0.4", "REST backend and application services; measured in the project virtual environment"),
    ("Django REST Framework", "3.17.1", "API serialization, view sets and permissions; measured in the project virtual environment"),
    ("PostgreSQL", "18.3", "Development relational database; server and client versions measured locally"),
    ("Node.js", "25.8.1", "Frontend build tooling and report-service runtime; measured locally"),
    ("React / React DOM", "19.2.8", "Single-page web dashboard; exact package-lock.json resolution"),
    ("Vite", "8.2.2", "Frontend development server and production bundler; exact lockfile resolution"),
    ("Tailwind CSS", "4.3.3", "Dashboard styling; exact lockfile resolution"),
    ("Ultralytics", "8.4.41", "YOLO training, validation and inference; measured in ai-module/.venv"),
    ("Django Channels", "4.3.2", "WebSocket alert delivery; measured in the project virtual environment"),
    ("Daphne", "4.2.3", "ASGI server; measured in the project virtual environment"),
    ("Git", "2.53.0", "Version control; measured locally"),
    ("Docker", "Dockerfile only", "Report-service image definition exists; deployment and orchestration are not verified"),
    ("MQTT / Mosquitto", "Scaffolding only", "Listener code exists but is not started; no broker deployment is present or host-verified"),
]

TABLE_47 = [
    ("Endpoint", "Method", "Description"),
    ("/api/token/", "POST", "Create a JWT access/refresh token pair"),
    ("/api/token/refresh/", "POST", "Refresh a JWT access token"),
    ("/api/dashboard/stats/", "GET", "Dashboard worker and alert counts"),
    ("/api/alerts/", "GET/POST", "List or create alerts"),
    ("/api/alerts/{id}/", "GET/PUT/PATCH", "Read or update one alert"),
    ("/api/alerts/{id}/resolve/", "POST", "Resolve one alert"),
    ("/api/alerts/simulate/", "POST", "Create a simulated alert"),
    ("/api/alerts/ingest/", "POST", "Ingest an alert from an edge/AI source"),
    ("/api/workers/", "GET/POST", "List or register workers"),
    ("/api/workers/{id}/", "GET/PUT/PATCH/DELETE", "Read, update or delete one worker"),
    ("/api/v1/vests/", "GET/POST", "List or register Smart Vest records"),
    ("/api/v1/vests/{id}/", "GET/PUT/PATCH/DELETE", "Read, update or delete one vest record"),
    ("/api/v1/vests/telemetry/", "POST", "Accept vest GPS, battery and SOS telemetry"),
    ("/api/v1/activity/", "GET/POST", "List or create activity records"),
    ("/api/v1/activity/{id}/", "GET", "Read one activity record"),
    ("/api/v1/activity/summary/?worker_id={id}&date=YYYY-MM-DD", "GET", "Summarize one worker's activity for a date"),
    ("/api/v1/detections/", "GET/POST", "List or create detection records"),
    ("/api/v1/detections/{id}/", "GET", "Read one detection record"),
    ("/api/v1/alert-config/", "GET/POST", "List or create alert configurations"),
    ("/api/v1/alert-config/{id}/", "GET/PUT/PATCH/DELETE", "Read, update or delete one alert configuration"),
    ("/api/v1/zones/", "GET/POST", "List or create zones"),
    ("/api/v1/zones/{id}/", "GET/PUT/PATCH/DELETE", "Read, update or delete one zone"),
    ("/api/v1/zones/{id}/workers/", "GET", "List workers assigned to a zone"),
    ("/api/v1/zones/{id}/check_point/", "POST", "Test a latitude/longitude against a zone boundary"),
    ("/api/v1/fcm-tokens/", "POST", "Register an FCM device token"),
    ("/api/v1/fcm-tokens/{token}/", "DELETE", "Remove an FCM device token"),
    ("/api/v1/profile/", "GET/PATCH", "Read or update the signed-in user's profile"),
    ("/api/v1/profile/change-password/", "POST", "Change the signed-in user's password"),
    ("/api/v1/reports/", "GET", "List generated reports"),
    ("/api/v1/reports/generate/", "POST", "Generate and store a DOCX report"),
    ("/api/v1/reports/{id}/download/", "GET", "Download one generated report"),
]

TABLE_48 = [
    ("Table", "Key fields", "Relations / notes"),
    ("users", "id, email, full_name, role, is_admin, is_active", "One-to-many: fcm_tokens, reports; one-to-one: user_preferences"),
    ("workers", "id, name, vest_id, phone, role, zone_id, is_active", "FK: zones; one-to-one: smart_vests; one-to-many: activity_logs, alerts, detections"),
    ("smart_vests", "id, vest_id, firmware_version, battery_level, latitude, longitude, sos_active, is_online, last_seen, worker_id", "One-to-one: workers; GPS/SOS live here; no vitals table"),
    ("detections", "id, camera_id, class_name, confidence, bbox, frame_timestamp, worker_id", "FK: workers; camera_id is text, not a camera-table FK"),
    ("alerts", "id, alert_type, severity, source, worker_id, camera_id, location, is_resolved, timestamp", "FK: workers; camera_id is text, not a camera-table FK"),
    ("zones", "id, name, risk_level, boundaries, camera_ids, max_occupancy, is_active", "One-to-many: workers, alert_config; camera_ids is JSON"),
    ("activity_logs", "id, worker_id, zone_id, action, timestamp, metadata", "FK: workers, zones; one timestamp per action"),
    ("alert_config", "id, zone_id, alert_type, is_enabled, threshold_seconds, notify_email, notify_push", "FK: zones"),
    ("fcm_tokens", "id, user_id, token, device_name, is_active, created_at", "FK: users"),
    ("reports", "id, report_type, date_from, date_to, file, created_by_id, created_at", "FK: users"),
    ("user_preferences", "id, user_id, notify_push, notify_email", "One-to-one: users"),
]

# These values are patched only after the implementation task's focused and
# full-suite tests have completed. The build command refuses placeholders.
#
# Reproduce: cd backend && ../.venv/bin/python manage.py test \
#              --settings=sitesafe.test_settings
# 35 tests, 0 failures, measured against the isolated SQLite test settings
# added alongside this work (production/development remain PostgreSQL).
AUTOMATED_TEST_ROWS: list[tuple[str, str, str]] = [
    (
        "Role-based access control (16 cases)",
        "Per-role API behavior: authenticated reads, alert resolution, worker "
        "creation, zone create/delete, and role exposure on the profile "
        "endpoint, including the legacy is_admin fallback",
        "Pass",
    ),
    (
        "Vehicle proximity logic (9 cases)",
        "Bounding-box gap geometry, scale invariance across box size, the "
        "persistence filter, re-arm after the condition clears, and "
        "detection-class splitting",
        "Pass",
    ),
    (
        "Camera-level inactivity tracking (10 cases)",
        "Centroid tracking, multi-person association under reordering, "
        "timeout and movement-reset behavior, alert latch and re-arm, "
        "tolerance for brief missed detections, stale-track eviction, and "
        "the alert API integration",
        "Pass",
    ),
]
AUTOMATED_TEST_SUMMARY = (
    "Automated coverage totals 35 tests across three areas, all passing at the "
    "time of writing. Table 5.1 groups them by subject rather than listing "
    "every case individually. Role-based access control tests exercise the "
    "Django REST Framework test client against the live URL configuration, "
    "so a permission check that regresses in routing or view configuration "
    "fails the same test that checks the underlying logic. The proximity "
    "and inactivity tests exercise ai-module/proximity.py and "
    "ai-module/inactivity.py directly; both modules are deliberately free of "
    "OpenCV, torch and Ultralytics imports, which is what allows their "
    "decision logic to run under the Django test command without a camera, a "
    "GPU or a loaded model. The suite runs against an isolated SQLite "
    "database created for this purpose so it does not require PostgreSQL "
    "CREATEDB permission or touch development data; one test that resolves "
    "a camera to a zone mocks that single lookup, noted in its own comment, "
    "because SQLite does not implement the JSON containment operator "
    "PostgreSQL provides for that query."
)
INACTIVITY_IMPLEMENTATION_PARAGRAPHS = (
    "Camera-level inactivity tracking is implemented in "
    "ai-module/inactivity.py and integrated into the main inference loop "
    "alongside PPE and proximity detection, sharing the same sampled "
    "person-detector frames used for proximity. Detected people are "
    "associated between sampled frames by nearest centroid, each assigned a "
    "temporary track id; a track is treated as stationary once its centroid "
    "stays within a movement threshold of its anchor position for at least "
    "a configured timeout, and dropped after a small number of consecutive "
    "missed detections rather than on the first one, since brief occlusion "
    "is routine on a working site. This produces a camera-local track, not "
    "a worker identity: it cannot survive the camera losing and "
    "re-acquiring a person, and pixel displacement is not a calibrated "
    "real-world distance.",
    "A confirmed track raises exactly one INACTIVITY alert for its "
    "stationary period; the track re-arms only once its displacement "
    "exceeds the movement threshold, so a worker who remains still does not "
    "generate a repeating stream of alerts. The alert is posted through the "
    "existing alert API with a camera_id and no worker, since the COCO "
    "detector supplying person boxes carries no worker identity; zone "
    "attribution therefore comes from the camera-to-zone association "
    "described in Section 4.3.6, not from the alert itself. Zone entry/exit "
    "logging, break-area detection and hourly activity_logs aggregation "
    "described in the original design are not implemented by this module.",
)


ABSTRACT_PARAGRAPHS = (
    "Construction sites need evidence that safety controls were followed, yet manual inspections and passive CCTV provide only intermittent oversight. Site-Safe is a camera-first safety compliance project that records PPE violations, routes alerts, and presents incident and workforce information through a web command centre. The report distinguishes the implemented software from proposed extensions so that evaluation is tied to repository evidence rather than design intent.",
    "The implemented system combines a ten-class YOLO11 PPE model, a separate pretrained person-and-vehicle detector for image-plane proximity warnings, configurable camera identifiers, camera-to-zone association, a Django REST and WebSocket backend, role-based access control, a React dashboard, and DOCX report generation. The backend also contains Smart Vest data models, HTTP telemetry endpoints and a dashboard simulator. QR identification, assembled Smart Vest/LoRa hardware, fall detection, fire/smoke detection and a React Native mobile application remain proposed designs and are not presented as delivered features.",
    "Held-out test evaluation measured mAP@0.5 of 0.544 and mAP@0.5:0.95 of 0.333. Present-equipment classes such as boots, vest and helmet exceeded 0.94 mAP@0.5, while no-gloves, no-boots and no-helmet measured 0.095, 0.175 and 0.308 respectively. This gap limits the product's core violation-detection purpose and makes balanced retraining the highest-value model improvement. Automated API and decision-logic tests and recorded end-to-end exercises are reported separately from manual or hardware verification that has not yet been completed.",
)


CHAPTER5_SECTIONS = {
    "5.1 Testing Approach": (
        "Verification is organized into three tiers. The first tier is automated software testing: API role enforcement is exercised through Django's test client, while the camera-side geometry and temporal decision modules are tested without loading OpenCV or a neural network. The second tier is model evaluation using Ultralytics on the held-out test split defined by the repository dataset. The third tier is manual end-to-end verification through the running API, dashboard simulator, synthetic camera sources and recorded detector exercises.",
        "The tiers answer different questions and are not interchangeable. Automated tests establish repeatable behavior for code paths covered by assertions, model validation measures detection quality on one labelled dataset, and manual exercises show that selected components can exchange real payloads. None of these alone proves reliability on an operating construction site, so each result below states its evidence and remaining verification boundary.",
    ),
    "5.2 Automated Test Results": (),
    "5.3 Model Evaluation Results": (
        "Table 4.6 reports the held-out test results from the currently deployed nano weights. The evaluation measured overall mAP@0.5 of 0.544, mAP@0.5:0.95 of 0.333, mean precision of 0.481 and mean recall of 0.665. These values describe the recorded model only; no result from the prepared medium-model retraining workflow is included because that training run has not yet been completed.",
        "The per-class values show a product-relevant asymmetry. Boots, vest and helmet reached 0.981, 0.956 and 0.942 mAP@0.5, while their absence classes no-boots, no-vest and no-helmet reached 0.175, 0.737 and 0.308. Gloves and no-gloves measured 0.304 and 0.095, and goggles and no-goggles measured 0.526 and 0.414. The model is therefore more reliable at recognizing visible equipment than at identifying several missing-equipment conditions. Because missing PPE is the alert condition, the aggregate score overstates usefulness for the core task unless the negative classes are inspected separately.",
        "The result is limited to one labelled test split from the same dataset family as training. It has not been cross-validated against a second site, different camera geometry, night-time conditions or a physical NVR feed. Those gaps prevent a claim of field generalization even though the stored validation output is reproducible.",
    ),
    "5.4 End-to-End Verification": (
        "Table 5.2 separates recorded manual outcomes from automated evidence. The retained results show that selected payloads and UI paths worked in prior exercises, but the detector images, synthetic stream harness and UI session are not stored as automated regression fixtures. They are therefore labelled as recorded or partial rather than treated as continuously verified tests.",
    ),
    "5.5 Discussion": (
        "The automated results support two narrow conclusions: the tested role boundaries are enforced by the API, and the tested geometry and temporal filters behave deterministically for their asserted inputs. They do not establish complete backend correctness, production security or physical-site performance. In particular, alert ingestion, detection creation and vest telemetry currently permit unauthenticated writes; the vision script contains hard-coded development credentials; and Channels uses an in-memory layer. These are acceptable prototype constraints only when stated explicitly and must be hardened before deployment.",
        "The model evaluation exposes the most important technical weakness. Several absence classes perform substantially worse than their presence counterparts, so the deployed weights can miss the very violations the system is intended to surface. A longer medium-model run at 640 pixels is prepared as future work, but improvement must be established by a new held-out evaluation rather than assumed from model size or training duration.",
        "Threats to validity remain in the evaluation setup. Automated API tests use an isolated SQLite database and therefore do not cover PostgreSQL-specific behavior. Recorded end-to-end exercises rely on simulated alerts, a synthetic stream and still images rather than a physical NVR, assembled vest or live industrial traffic. Dashboard trend charts also include seeded demonstration series rather than a fully historical analytics pipeline. The evidence supports an FYP prototype and identifies its boundaries; it does not support a production-readiness claim.",
    ),
}


TABLE_52 = [
    ("Verification scenario", "Expected behavior", "Recorded status and boundary"),
    ("Proximity with the pretrained detector (still image)", "Resolve person and vehicle boxes and classify breach severity", "Prior recorded pass: four person boxes and one vehicle box; CRITICAL event. Not rerun as an automated fixture in this revision."),
    ("Distant-vehicle negative case", "Keep the alert silent while separation exceeds the configured threshold", "Prior recorded pass across repeated frames; not retained as an automated detector fixture."),
    ("Alert with a camera identifier", "Store the alert and associate its camera with a configured zone", "Prior recorded pass: HTTP 201 and zone resolved from camera_id; not rerun as a browser workflow in this revision."),
    ("Dashboard alert simulation", "Create an alert and show it in the live feed and incident log", "Prior recorded pass: HTTP 201 and both views updated; manual UI session not rerun in this revision."),
    ("Camera-source and read-retry path", "Accept device indices and URLs and tolerate bounded dropped reads", "Partial: verified with a synthetic source; not verified against a physical NVR or phone stream."),
]


def normalized(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").replace("–", "-").replace("—", "-").split())


def paragraph_text(element) -> str:
    return "".join(element.itertext()).strip()


def find_paragraph(doc: Document, text: str, *, exact: bool = True):
    matches = [
        p for p in doc.paragraphs
        if (p.text.strip() == text if exact else p.text.strip().startswith(text))
    ]
    if len(matches) != 1:
        raise AssertionError(f"Expected one paragraph for {text!r}; found {len(matches)}")
    return matches[0]


def clear_paragraph(p) -> None:
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def set_paragraph_text(p, text: str) -> None:
    clear_paragraph(p)
    p.add_run(text)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clone_paragraph(template, text: str | None = None):
    element = deepcopy(template._p)
    if text is not None:
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        run = OxmlElement("w:r")
        txt = OxmlElement("w:t")
        txt.set(qn("xml:space"), "preserve")
        txt.text = text
        run.append(txt)
        element.append(run)
    return element


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_table_geometry(table: Table, widths_inches: list[float]) -> None:
    widths = [int(Inches(value)) for value in widths_inches]
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        for height in tr_pr.findall(qn("w:trHeight")):
            tr_pr.remove(height)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[min(idx, len(widths) - 1)]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text: str, *, header: bool = False, size: float = 9.0) -> None:
    p = cell.paragraphs[0]
    set_paragraph_text(p, text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    for run in p.runs:
        set_run_font(run, size=size, bold=header)


def replace_table_rows(table: Table, rows: list[tuple[str, ...]], widths: list[float], *, font_size: float = 9.0) -> None:
    if len(rows[0]) != len(table.columns):
        raise AssertionError("Replacement column count does not match table")
    template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for values in rows[1:]:
        table._tbl.append(deepcopy(template))
        row = table.rows[-1]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, size=font_size)
    for cell, value in zip(table.rows[0].cells, rows[0]):
        set_cell_text(cell, value, header=True, size=font_size)
    set_repeat_header(table.rows[0])
    set_table_geometry(table, widths)


def table_after_caption(doc: Document, caption: str) -> Table:
    p = find_paragraph(doc, caption)
    next_el = p._p.getnext()
    if next_el is None or next_el.tag != qn("w:tbl"):
        raise AssertionError(f"No table follows caption {caption!r}")
    for table in doc.tables:
        if table._tbl is next_el:
            return table
    raise AssertionError(f"Table wrapper not found after {caption!r}")


def clone_table_from_xml(doc: Document, xml) -> Table:
    element = deepcopy(xml)
    return Table(element, doc._body)


def insert_before(reference_element, elements) -> None:
    for element in elements:
        reference_element.addprevious(element)


def make_index_paragraph(text: str, page: int, level: int) -> OxmlElement:
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "8500")
    tabs.append(tab)
    p_pr.append(tabs)
    if level > 1:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str((level - 1) * 300))
        p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "20")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    p.append(p_pr)

    def add_run(value: str, *, bold: bool = False) -> None:
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "cs"):
            fonts.set(qn(f"w:{attr}"), "Times New Roman")
        r_pr.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "20")
        r_pr.append(size)
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), "20")
        r_pr.append(size_cs)
        if bold:
            r_pr.append(OxmlElement("w:b"))
        run.append(r_pr)
        txt = OxmlElement("w:t")
        txt.set(qn("xml:space"), "preserve")
        txt.text = value
        run.append(txt)
        p.append(run)

    add_run(text, bold=level == 1)
    p.append(OxmlElement("w:r"))
    p[-1].append(OxmlElement("w:tab"))
    add_run(str(page))
    return p


def replace_region(doc: Document, start_title: str, end_title: str, entries: list[tuple[str, int, int]]) -> None:
    start = find_paragraph(doc, start_title)._p
    end = find_paragraph(doc, end_title)._p
    node = start.getnext()
    while node is not None and node is not end:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node
    for text, page, level in entries:
        end.addprevious(make_index_paragraph(text, page, level))


def body_start_index(doc: Document) -> int:
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Chapter 1":
            return idx
    raise AssertionError("Chapter 1 marker not found")


def collect_toc_entries(doc: Document) -> list[dict]:
    start = body_start_index(doc)
    paragraphs = doc.paragraphs[start:]
    entries: list[dict] = []
    for idx, p in enumerate(paragraphs):
        text = p.text.strip()
        match = re.fullmatch(r"Chapter ([1-6])", text)
        if match:
            chapter = int(match.group(1))
            title = CHAPTER_TITLES[chapter]
            entries.append({"display": f"Chapter {chapter}: {title}", "search": title, "level": 1})
            continue
        if text == "References":
            entries.append({"display": "References", "search": "References", "level": 1})
            break
        style = p.style.name if p.style else ""
        if text and style == "Heading 2":
            entries.append({"display": text, "search": text, "level": 2})
        elif text and style == "Heading 3":
            entries.append({"display": text, "search": text, "level": 3})
    return entries


def collect_captions(doc: Document, kind: str) -> list[str]:
    start = body_start_index(doc)
    pattern = re.compile(rf"^{kind} \d+\.\d+:\s+.+")
    return [normalized(p.text) for p in doc.paragraphs[start:] if pattern.match(normalized(p.text))]


def write_front_matter(doc: Document, *, page_lookup: dict[str, int] | None = None) -> None:
    toc = collect_toc_entries(doc)
    figures = collect_captions(doc, "Figure")
    tables = collect_captions(doc, "Table")
    page_lookup = page_lookup or {}
    replace_region(
        doc,
        "TABLE OF CONTENTS",
        "LIST OF FIGURES",
        [(entry["display"], page_lookup.get(entry["search"], 0), entry["level"]) for entry in toc],
    )
    replace_region(
        doc,
        "LIST OF FIGURES",
        "LIST OF TABLES",
        [(caption, page_lookup.get(caption, 0), 1) for caption in figures],
    )
    replace_region(
        doc,
        "LIST OF TABLES",
        "LIST OF ABBREVIATIONS",
        [(caption, page_lookup.get(caption, 0), 1) for caption in tables],
    )


def replace_abstract(doc: Document) -> None:
    current = find_paragraph(doc, "Construction sites in Pakistan account", exact=False)
    set_paragraph_text(current, ABSTRACT_PARAGRAPHS[0])
    reference = current._p.getnext()
    insert_before(
        reference,
        [clone_paragraph(current, ABSTRACT_PARAGRAPHS[1]), clone_paragraph(current, ABSTRACT_PARAGRAPHS[2])],
    )


def replace_chapter4_content(doc: Document) -> None:
    replacements = {
        "Our development workflow spans several machines and platforms": "The verified development environment combines a Django REST backend, PostgreSQL, a React/Vite dashboard, an Ultralytics vision module and a Node report service. Versions in Table 4.1 were measured from the active virtual environments, local executables or exact package-lock.json resolutions. Unpinned editor and API-client versions are omitted. Recorded model runs used local CPU or MPS settings; a new medium-model GPU run is prepared but has not been completed.",
        "4.2 QR Identification and Smart Vest Design": "4.2 Proposed QR Identification and Smart Vest Design",
        "The Smart Vest is a lightweight supplementary wearable designed": "The Smart Vest and QR artifacts in this section are a proposed hardware and identification design, not an assembled implementation. The repository contains a SmartVest database model, HTTP telemetry endpoint and dashboard simulator for GPS, battery and SOS payloads. It does not contain ESP32 firmware, a LoRa link, a QR field on Worker, QR generation or a camera-side QR decoder.",
        "The hardware simulation panel in the web dashboard mirrors": "The simulator exercises selected backend and web-dashboard paths without physical hardware. It does not verify LoRa transport, a real GPS receiver, an SOS button, a native mobile client or field behavior. Figure 4.1 and Tables 4.2 and 4.3 therefore document the proposed component design and pin allocation only.",
        "Figure 4.1: Smart Vest Hardware Implementation": "Figure 4.1: Proposed Smart Vest Hardware Design",
        "Table 4.2: Smart Vest Component Specifications": "Table 4.2: Proposed Smart Vest Component Specifications",
        "Table 4.3: ESP32 Pin Allocation": "Table 4.3: Proposed ESP32 Pin Allocation",
        "Our computer vision pipeline uses a YOLOv11-M model": "The deployed PPE pipeline uses the nano YOLO11 weights stored at ai-module/models/sitesafe_final.pt. It was trained on a ten-class Roboflow-format PPE dataset. Vehicle and person classes are absent from that PPE model, so a separate pretrained YOLO11 nano detector supplies person and vehicle boxes for proximity and activity logic.",
        "4.3.4 Fall Detection via Pose Estimation": "4.3.4 Proposed Fall Detection via Pose Estimation (Not Implemented)",
        "Fall detection is implemented using YOLOv11n-Pose": "The proposed fall-detection extension would use a lightweight YOLO11 pose model to obtain body keypoints. No pose weights are loaded by the current inference pipeline, FALL is not an AlertType value, and no fall-detection module or migration exists in the repository.",
        "The fall detection algorithm computes two geometric features": "The design would compare horizontal and vertical body spans and hip-to-knee geometry, then require a sustained posture before raising an event. These thresholds are design parameters only and have not been evaluated on this project dataset.",
        "Both conditions must be met simultaneously": "A production implementation would need a testable pose module, persistence filtering, labelled fall/non-fall footage and a newly migrated FALL alert type. No accuracy range or alert behavior is claimed for the current system because that implementation and validation have not occurred.",
        "4.3.5 Fire and Smoke Detection with Zone-Aware Suppression": "4.3.5 Proposed Fire and Smoke Detection with Zone-Aware Suppression (Not Implemented)",
        "Fire and smoke detection uses a YOLO model": "Fire and smoke detection remains a proposed extension. The repository has no fire/smoke dataset, trained weights, detector invocation or FIRE/SMOKE AlertType values. The intended design would use a separately trained detector rather than the ten-class PPE model.",
        "The primary engineering challenge is avoiding false alarms": "The proposed zone-aware layer would distinguish scheduled hot work from an uncontrolled event by consulting zone context and permit data. Neither a permitted_activities field nor a hot-work permit model exists in the implemented schema, so suppression is not currently available.",
        "A growth-rate heuristic provides a second line of defence": "A temporal growth heuristic is a design hypothesis that would require labelled validation before thresholds or severity changes could be justified. It is retained as future work and is not part of the current alert pipeline.",
        "The web dashboard is a React single-page application": "The web dashboard is a React single-page application built with Vite. It uses REST endpoints for data and a WebSocket connection for live alerts. Seven implemented routes cover the dashboard, workers, incidents, vest records, zones, reports and settings. Recharts renders the current dashboard charts; the hourly and compliance trend series include seeded demonstration values rather than a complete historical analytics feed.",
        "The dashboard layout follows a sidebar navigation pattern": "Role information is exposed by the backend and server-side permissions enforce Admin, Safety Officer and Viewer operations. The interface presents summary counts, live alerts, worker and vest management, incident resolution, zones, generated reports and configuration. It does not contain the proposed native mobile screens, QR identity workflow or a GPS map.",
        "4.6 Mobile Application Implementation": "4.6 Proposed Mobile Application Design (Not Implemented)",
        # NB: two paragraphs share the opening clause "The mobile application
        # is built with React Native" — this one (§4.6.1, "...chosen for
        # cross-platform support...") is the Chapter 4 implementation claim
        # being corrected. The other, in §3.6.1 ("...enabling a single
        # JavaScript codebase..."), is Chapter 3 proposed-design text and is
        # intentionally left untouched.
        "The mobile application is built with React Native, chosen for cross-platform support": "The React Native client is a proposed design only. No mobile source tree, dependency manifest or compiled application exists in the repository. The web dashboard is the sole implemented user interface.",
        "4.6.1 Technology Stack": "4.6.1 Proposed Technology Stack",
        "Table 4.9: Mobile App Technology Stack": "Table 4.9: Proposed Mobile Application Technology Stack",
        "4.6.2 Core Screens": "4.6.2 Proposed Core Screens",
        "The mobile app has five primary screens": "The wireframes specify login, alert feed, worker status, activity summary and alert-configuration screens for a possible future client. They describe intended scope and are not evidence of implemented navigation, biometric unlock, QR search or activity timelines.",
        "4.6.3 Push Notification Pipeline": "4.6.3 Proposed Push Notification Flow",
        "Push notifications follow a pipeline": "The backend stores FCM tokens and contains notification integration points, but no native client is available to receive and display those messages. The end-to-end mobile push flow remains unverified and is retained as a proposed extension.",
        "4.7 Workforce Activity Monitoring Implementation": "4.7 Workforce Activity Monitoring (Partial Implementation)",
        "Figure 4.4: Workforce Activity Monitoring Pipeline": "Figure 4.4: Workforce Activity Monitoring Design",
        "4.8 Data Flow": "4.8 Implemented and Proposed Data Flow",
        "Figure 4.5: System Data Flow Diagram": "Figure 4.5: Implemented and Proposed System Data Flow",
        "The data flow diagram illustrates how information moves through the system": "Figure 4.5 combines implemented and proposed paths. The implemented camera path posts PPE and proximity alerts with camera identifiers to the Django API, which stores and broadcasts them to the web dashboard. HTTP vest telemetry, report generation and alert simulation are also implemented. QR decoding, LoRa/MQTT runtime ingestion, a native mobile client and mobile push display remain proposed; the MQTT listener exists as unstarted scaffolding.",
        "4.9 SOS and Location Alert Flow": "4.9 Proposed SOS and Location Alert Flow",
        "Figure 4.6: SOS and Location Alert Flowchart": "Figure 4.6: Proposed SOS and Location Alert Flow",
        "The SOS and location alert flow handles the Smart Vest’s two primary functions": "Figure 4.6 describes the proposed physical SOS and location flow. The current backend can accept simulated or HTTP telemetry and can store location/SOS state, but no assembled vest, LoRa transport or map visualization has been verified. The figure must therefore be read as the target design rather than an implemented hardware sequence.",
        "4.10 End-to-End Alert Flow": "4.10 Proposed End-to-End Alert Flow",
        "Figure 4.7: End-to-End Alert Processing Flow": "Figure 4.7: Proposed End-to-End Alert Processing Flow",
        "The diagram above traces an alert from origin to resolution": "Figure 4.7 is a composite target flow. The implemented path accepts PPE, vehicle-proximity, SOS, zone-breach and inactivity enum values, stores alerts, broadcasts WebSocket updates, and supports web-dashboard resolution. Automatic fall, fire/smoke, QR identity, hot-work suppression, native mobile display and assembled Smart Vest inputs are not implemented. Those nodes remain in the diagram to show the intended architecture, not the delivered result.",
        "The project is divided into 13 major tasks": "The project plan contains research, development, testing and documentation tasks. Mobile application, Smart Vest hardware, QR identification, fall detection and fire/smoke detection are design or future-work tasks rather than completed implementation. Verification status is reported in Chapter 5 and deferred work in Section 6.2.",
    }
    for prefix, value in replacements.items():
        p = find_paragraph(doc, prefix, exact=prefix[0].isdigit() or prefix.startswith(("Figure", "Table")))
        set_paragraph_text(p, value)

    if len(INACTIVITY_IMPLEMENTATION_PARAGRAPHS) != 2:
        raise RuntimeError("Set the two verified inactivity implementation paragraphs before building")
    set_paragraph_text(
        find_paragraph(doc, "The workforce activity monitoring module adds", exact=False),
        INACTIVITY_IMPLEMENTATION_PARAGRAPHS[0],
    )
    set_paragraph_text(
        find_paragraph(doc, "The inactivity detection works by maintaining", exact=False),
        INACTIVITY_IMPLEMENTATION_PARAGRAPHS[1],
    )

    # Repair the sole figure-number mismatch: the body said 3.9 while the list
    # and sequence require 3.8.
    set_paragraph_text(
        find_paragraph(doc, "Figure 3.9: Mobile Application Screen Wireframes"),
        "Figure 3.8: Mobile Application Screen Wireframes",
    )

    # Repair two pre-existing Chapter 1/3 numbering defects, found by
    # inspecting the rebuilt table of contents against the source document:
    # 1.5 Significance carries body ("Normal") style rather than Heading 2,
    # so it silently drops out of any TOC built from heading styles; and the
    # heading text "3.3 Architecture" duplicates the number already used two
    # headings later by "3.3 Module 1: Vision Intelligence", leaving "3.2"
    # unused except as the orphaned parent of subsections 3.2.1/3.2.2.
    significance = find_paragraph(doc, "1.5 Significance")
    significance.style = doc.styles["Heading 2"]
    set_paragraph_text(
        find_paragraph(doc, "3.3 Architecture"),
        "3.2 Architecture",
    )


def restructure_chapters(doc: Document) -> None:
    if not AUTOMATED_TEST_ROWS or not AUTOMATED_TEST_SUMMARY:
        raise RuntimeError("Set verified automated-test rows and summary before building")

    integration_heading = find_paragraph(doc, "4.11 Integration Testing")
    integration_prose = find_paragraph(doc, "Verification to date falls", exact=False)
    old_caption = find_paragraph(doc, "Table 4.10: Integration Test Scenarios and Results")
    old_table = table_after_caption(doc, "Table 4.10: Integration Test Scenarios and Results")
    table_template_xml = deepcopy(old_table._tbl)
    for element in (integration_heading._p, integration_prose._p, old_caption._p, old_table._tbl):
        element.getparent().remove(element)

    timeline = find_paragraph(doc, "4.12 Project Timeline")
    set_paragraph_text(timeline, "4.11 Project Timeline")

    old_chapter = find_paragraph(doc, "Chapter 5")
    old_title = find_paragraph(doc, "Conclusion and Future Work")
    old_51 = find_paragraph(doc, "5.1 Conclusion")
    old_52 = find_paragraph(doc, "5.2 Future Work")
    page_break = old_chapter._p.getprevious()
    if page_break is None:
        raise AssertionError("Chapter 5 page break not found")

    body_template = find_paragraph(doc, "This project set out to build", exact=False)
    heading_template = old_51
    caption_template = old_caption

    elements = [
        clone_paragraph(old_chapter, "Chapter 5"),
        clone_paragraph(old_title, CHAPTER_TITLES[5]),
    ]

    for heading, paragraphs in CHAPTER5_SECTIONS.items():
        elements.append(clone_paragraph(heading_template, heading))
        if heading == "5.2 Automated Test Results":
            elements.append(clone_paragraph(body_template, AUTOMATED_TEST_SUMMARY))
            elements.append(clone_paragraph(caption_template, "Table 5.1: Automated Test Results"))
            table_51 = clone_table_from_xml(doc, table_template_xml)
            replace_table_rows(
                table_51,
                [("Test group", "Verified scope", "Result"), *AUTOMATED_TEST_ROWS],
                [1.65, 2.95, 1.2],
                font_size=8.5,
            )
            elements.append(table_51._tbl)
        else:
            for paragraph in paragraphs:
                elements.append(clone_paragraph(body_template, paragraph))
        if heading == "5.4 End-to-End Verification":
            elements.append(clone_paragraph(caption_template, "Table 5.2: End-to-End Verification Record"))
            table_52 = clone_table_from_xml(doc, table_template_xml)
            replace_table_rows(table_52, TABLE_52, [1.65, 2.15, 2.2], font_size=8.5)
            elements.append(table_52._tbl)

    elements.append(deepcopy(page_break))
    insert_before(old_chapter._p, elements)

    set_paragraph_text(old_chapter, "Chapter 6")
    set_paragraph_text(old_51, "6.1 Conclusion")
    set_paragraph_text(old_52, "6.2 Future Work")

    conclusion_paragraphs = [
        "This project delivers a camera-first safety prototype with a ten-class PPE model, multi-source camera configuration, camera-to-zone association, vehicle-proximity logic, a Django REST and WebSocket backend, role-based access control, a React command centre and DOCX report generation. Smart Vest records and HTTP telemetry simulation supplement the camera path, while assembled hardware and native mobile software remain outside the delivered implementation.",
        "The findings qualify that achievement. The current model measured 0.544 mAP@0.5 on the held-out test split, but several missing-equipment classes scored far below the corresponding present-equipment classes. Automated tests verify the asserted authorization and camera-side decision logic, while end-to-end evidence is limited to recorded simulator, still-image and synthetic-stream exercises. The system therefore demonstrates an integrated FYP prototype, not field-validated production safety equipment.",
        "Against the original objectives, continuous PPE alerting, a web command centre, auditable alert records, document reports and configurable multi-camera attribution are implemented. QR worker identification, automatic fall and fire/smoke detection, physical vest/LoRa telemetry, a native mobile application and full worker-level activity accounting remain design or future-work items. This separation between results and proposals is central to interpreting the project honestly.",
    ]
    existing_conclusion = [
        find_paragraph(doc, "This project set out to build", exact=False),
        find_paragraph(doc, "Two objectives were met", exact=False),
        find_paragraph(doc, "The most significant finding", exact=False),
    ]
    for p, text in zip(existing_conclusion, conclusion_paragraphs):
        set_paragraph_text(p, text)

    future_paragraphs = [
        "The first priority is to improve and re-measure the PPE model. A reproducible Google Colab workflow is prepared for YOLO11m training at 640 pixels, 150 epochs with early stopping, AdamW and a 0.001 initial learning rate. No result from that run is included yet. The new best weights must be evaluated on the held-out test split and Table 4.5, Table 4.6 and Section 5.3 updated only with measured output. If negative classes remain weak, training-image sampling should be adjusted using the measured label distribution.",
        "QR identification, fall detection and fire/smoke detection require new code, data and validation. QR work needs a Worker field, QR generation and camera decoding; fall detection needs a tested pose module and FALL migration; fire/smoke needs a suitable labelled dataset, trained weights and explicit hot-work data models. The React Native application and assembled Smart Vest/LoRa hardware are deferred because neither has a repository artefact or physical verification.",
        "The implemented camera-level activity logic still needs identity-aware tracking, zone semantics, break classification and durable ActivityLog integration before it can support per-worker attendance claims. Physical verification is also outstanding for NVR/phone streams, live industrial vehicles and vest hardware. Production hardening must remove hard-coded credentials, authenticate ingestion endpoints, replace the in-memory channel layer and exercise PostgreSQL-specific behavior in an isolated database.",
    ]
    existing_future = [
        find_paragraph(doc, "Several capabilities described", exact=False),
        find_paragraph(doc, "Improving the vision model", exact=False),
        find_paragraph(doc, "Two pieces of verification", exact=False),
    ]
    for p, text in zip(existing_future, future_paragraphs):
        set_paragraph_text(p, text)

    # Literal section references: one remains in 4.3.3 after the testing block
    # is moved; all future-work references now point to Chapter 6.
    for p in doc.paragraphs:
        if "Section 5.2" in p.text:
            set_paragraph_text(p, p.text.replace("Section 5.2", "Section 6.2"))


def update_tables_and_references(doc: Document) -> None:
    replace_table_rows(table_after_caption(doc, "Table 4.1: Development Environment Summary"), TABLE_41, [1.45, 1.15, 3.4], font_size=8.5)
    replace_table_rows(table_after_caption(doc, "Table 4.7: Core API Endpoints"), TABLE_47, [2.8, 1.1, 2.1], font_size=7.8)
    replace_table_rows(table_after_caption(doc, "Table 4.8: Database Schema Summary"), TABLE_48, [1.15, 3.25, 1.6], font_size=7.8)

    set_paragraph_text(
        find_paragraph(doc, '[3] G. Jocher', exact=False),
        '[3] G. Jocher, J. Qiu, and A. Chaurasia, "Ultralytics YOLO," version 8.4.41, 2026. [Online]. Available: https://github.com/ultralytics/ultralytics',
    )
    set_paragraph_text(
        find_paragraph(doc, '[10] Django Software Foundation', exact=False),
        '[10] Django Software Foundation, "Django REST Framework," version 3.17.1, 2026. [Online]. Available: https://www.django-rest-framework.org/',
    )


def normalize_typography(doc: Document) -> None:
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Times New Roman")
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 3"].font.size = Pt(12)

    chapter_titles = set(CHAPTER_TITLES.values())
    front_titles = {"ABSTRACT", "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS"}
    chapter1 = body_start_index(doc)
    for idx, p in enumerate(doc.paragraphs):
        text = normalized(p.text)
        style = p.style.name if p.style else ""
        is_caption = bool(re.match(r"^(Table|Figure) \d+\.\d+:", text))
        is_separator = bool(re.fullmatch(r"Chapter [1-6]", text)) or text in chapter_titles
        for run in p.runs:
            if style.startswith("Heading"):
                set_run_font(run)
                run.font.size = None
            elif is_caption:
                set_run_font(run, size=10)
            elif idx >= chapter1 and not is_separator:
                set_run_font(run, size=12)
            elif text.startswith("Construction sites need evidence") or text in ABSTRACT_PARAGRAPHS:
                set_run_font(run, size=12)
            elif text in front_titles:
                set_run_font(run, size=16, bold=True)
            else:
                set_run_font(run)
        if idx >= chapter1 and text and not is_caption and not is_separator and not style.startswith("Heading"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
        if is_caption:
            p.paragraph_format.keep_with_next = True

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run)


def build_report(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    replace_abstract(doc)
    replace_chapter4_content(doc)
    restructure_chapters(doc)
    update_tables_and_references(doc)
    normalize_typography(doc)
    write_front_matter(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def pdf_pages(pdf_path: Path) -> list[str]:
    reader = PdfReader(str(pdf_path))
    return [normalized(page.extract_text() or "") for page in reader.pages]


def find_page(pages: list[str], search: str, start: int) -> int:
    term = normalized(search)
    probes = [term, term[:80], term[:55]]
    for probe in probes:
        if len(probe) < 8:
            continue
        for idx in range(start, len(pages)):
            if probe in pages[idx]:
                return idx
    raise AssertionError(f"Could not locate rendered page for {search!r}")


def measured_page_lookup(doc: Document, pdf_path: Path) -> dict[str, int]:
    pages = pdf_pages(pdf_path)
    front_matter_markers = (
        "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS",
    )
    # The Table of Contents page itself contains the literal text
    # "Chapter 1: Introduction" as its first entry, which satisfies a plain
    # substring check just as well as the real chapter-separator page does.
    # Excluding known front-matter pages finds the actual separator instead.
    chapter1 = next(
        idx for idx, text in enumerate(pages)
        if "Chapter 1" in text and CHAPTER_TITLES[1] in text
        and not any(marker in text for marker in front_matter_markers)
    )
    lookup: dict[str, int] = {}

    cursor = chapter1
    for entry in collect_toc_entries(doc):
        cursor = find_page(pages, entry["search"], cursor)
        lookup[entry["search"]] = cursor - chapter1 + 1

    for captions in (collect_captions(doc, "Figure"), collect_captions(doc, "Table")):
        cursor = chapter1
        for caption in captions:
            cursor = find_page(pages, caption, cursor)
            lookup[caption] = cursor - chapter1 + 1
    return lookup


def finalize_report(input_path: Path, pdf_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    lookup = measured_page_lookup(doc, pdf_path)
    write_front_matter(doc, page_lookup=lookup)
    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build")
    build.add_argument("input", type=Path)
    build.add_argument("output", type=Path)
    finalize = sub.add_parser("finalize")
    finalize.add_argument("input", type=Path)
    finalize.add_argument("pdf", type=Path)
    finalize.add_argument("output", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.command == "build":
        build_report(args.input, args.output)
    else:
        finalize_report(args.input, args.pdf, args.output)


if __name__ == "__main__":
    main()
